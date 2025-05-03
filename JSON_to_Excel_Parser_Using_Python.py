import json
import pandas as pd
import openpyxl
from openpyxl.styles import Font, PatternFill, Border, Side, Alignment
import pdb

#pdb.set_trace() 
# Load JSON data from a file
try:
    with open('data.json', 'r') as file:
        data = json.load(file)  # Correct JSON structure
except json.JSONDecodeError as e:
    print(f"Error reading JSON file: {e}")
    exit()
except FileNotFoundError:
    print("File not found! Ensure 'data.json' exists in the correct location.")
    exit()


# List to store data for creating a DataFrame
rows = []
all_subjects = set()

# First pass to gather all unique subjects
for state in data['states']:
    for college in state['colleges']:
        for standard in college['standards']:
            for student in standard['students']:
                all_subjects.update(student['subjects'].keys())
#pdb.set_trace() 
#print(all_subjects)

# Extract data into rows
for state in data['states']:
    state_name = state['name']
    for college in state['colleges']:
        college_name = college['name']
        for standard in college['standards']:
            standard_name = standard['standard']
            for student in standard['students']:
                student_name = student['name']
                # Add subjects and marks dynamically
                row = {
                    'State': state_name,
                    'College': college_name,
                    'Standard': standard_name,
                    'Student': student_name
                }
                row.update(student['subjects'])
                rows.append(row)
#pdb.set_trace() 
#print(rows)

# Create a DataFrame from the rows, removing duplicates just in case
df = pd.DataFrame(rows).drop_duplicates()
#print(df)

# Create a new Excel workbook
workbook = openpyxl.Workbook()
sheet = workbook.active

# Define the heading
heading_text = "Exam Results for Different States"
heading_rows = 3
heading_cols = 5

# Merge cells for the heading, starting after 3 columns
sheet.merge_cells(start_row=1, start_column=4, end_row=heading_rows, end_column=heading_cols + 3)

# Set the heading text
heading_cell = sheet.cell(row=1, column=4, value=heading_text)

# Apply styles to the heading
heading_cell.font = Font(bold=True, size=14)
heading_fill = PatternFill(start_color="FFCC99", end_color="FFCC99", fill_type="solid")
heading_cell.fill = heading_fill

# Center align the heading text
heading_cell.alignment = Alignment(horizontal="center", vertical="center")

# Create a border for the heading
thin_border = Border(left=Side(style='thin'), right=Side(style='thin'), top=Side(style='thin'), bottom=Side(style='thin'))
for row in range(1, heading_rows + 1):
    for col in range(4, heading_cols + 4):  # Adjust column range to match merged cells
        cell = sheet.cell(row=row, column=col)
        cell.border = thin_border

# Initialize the starting row for data
row_num = heading_rows + 2

# Define fill colors for states
state_colors = ['FFFF99', 'CCFFCC', 'FFCCCC', 'CCFFFF', 'FFCC99', 'CCCCFF']
state_color_index = 0

# Write data from the DataFrame into the Excel sheet with formatting
state_groups = df.groupby(['State', 'College'])
for (state, college), group in state_groups:
    # Write the state name and apply background color
    state_cell = sheet.cell(row=row_num, column=1, value=f"State: {state}")
    state_fill = PatternFill(start_color=state_colors[state_color_index], end_color=state_colors[state_color_index], fill_type="solid")
    state_cell.fill = state_fill
    state_cell.font = Font(bold=True)  # Make the state cell bold
    state_color_index = (state_color_index + 1) % len(state_colors)
    row_num += 1
    
    # Write the college name and make it bold
    college_cell = sheet.cell(row=row_num, column=1, value=f"College: {college}")
    college_cell.font = Font(bold=True)  # Make the college cell bold
    row_num += 1

    # Set headers for subjects
    headers = ['Name of Student', 'Standard'] + sorted(all_subjects)
    for col_num, header in enumerate(headers, start=1):
        cell = sheet.cell(row=row_num, column=col_num, value=header)
        cell.font = Font(bold=True)
    row_num += 1

    # Write rows for each student
    for _, row in group.iterrows():
        sheet.cell(row=row_num, column=1, value=row['Student'])
        sheet.cell(row=row_num, column=2, value=row['Standard'])
        
        # Write subjects and marks with center alignment
        for col_num, subject in enumerate(sorted(all_subjects), start=3):
            marks = row.get(subject, '')
            cell = sheet.cell(row=row_num, column=col_num, value=marks)
            cell.alignment = Alignment(horizontal="center", vertical="center")  # Center align subject data
        
        row_num += 1
    
    row_num += 1  # Add an extra row to separate colleges

# Save the Excel workbook
workbook.save('parsed_data.xlsx')
print("Data has been written to 'parsed_data.xlsx'")
